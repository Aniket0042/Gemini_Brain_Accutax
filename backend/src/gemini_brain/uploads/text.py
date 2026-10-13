"""Text out of files, with no model: text PDFs, Word files, spreadsheets, CSV. And the image shrinking Bedrock needs.

Everything here reads files the upload checks (`validate.inspect`) already accepted, so the size, page and row
limits have been applied. Output is capped so a large file cannot flood the assistant's context.
"""
from __future__ import annotations

import csv
import io
import logging
import re
from typing import List, Tuple

from gemini_brain.uploads.extract import MAX_TEXT_CHARS

logger = logging.getLogger("gemini_brain.uploads.text")

#: A PDF page with fewer characters than this on average is a scan: its text is in the pixels.
MIN_CHARS_PER_PAGE = 40

_BILL_WORDS = re.compile(
    r"tax invoice|invoice|receipt|bill\b|vat\b|\btrn\b|tax registration|gst|total due|amount due|grand total|"
    r"statement of account|account statement|credit note|debit note|فاتورة|إيصال|ضريبة", re.I)

# Bedrock Converse: an image is at most 3.75 MB and 8,000 px on a side. Large photos also cost nothing extra
# above about 1,568 px (the model shrinks them), so keep detail for small print but stop at 3,000.
MAX_IMAGE_SIDE = 3_000
MAX_IMAGE_BYTES = 3_500_000


def pdf_pages(content: bytes) -> List[str]:
    """Text of each page, in layout order. A page with no text layer gives ''. [] when the PDF cannot be read."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text(extraction_mode="layout") or "")
            except Exception:
                pages.append("")
        return pages
    except Exception as e:
        logger.info("PDF text could not be read: %s", e)
        return []


def split_pdf(content: bytes, size: int) -> List[Tuple[int, int, bytes]]:
    """The PDF cut into parts of `size` pages: [(first page, last page, bytes)], pages counted from 1."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(content), strict=False)
    total = len(reader.pages)
    parts = []
    for start in range(0, total, size):
        writer = PdfWriter()
        for index in range(start, min(start + size, total)):
            writer.add_page(reader.pages[index])
        buf = io.BytesIO()
        writer.write(buf)
        parts.append((start + 1, min(start + size, total), buf.getvalue()))
    return parts


def is_scanned(pages: List[str]) -> bool:
    """No usable text layer: the pages are pictures."""
    if not pages:
        return True
    chars = sum(len(re.sub(r"\s", "", p)) for p in pages)
    return chars / len(pages) < MIN_CHARS_PER_PAGE


def looks_like_bill(text: str) -> bool:
    """Two or more bill or statement words: worth a structured read."""
    return len(set(m.group(0).lower() for m in _BILL_WORDS.finditer(text or ""))) >= 2


def join_pages(pages: List[str]) -> str:
    text = "\n\n".join(f"[Page {i}]\n{p.strip()}" for i, p in enumerate(pages, start=1) if p.strip())
    return text[:MAX_TEXT_CHARS]


def docx_text(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for number, table in enumerate(document.tables, start=1):
        rows = [" | ".join(c.text.strip().replace("\n", " ") for c in row.cells) for row in table.rows]
        parts.append(f"[Table {number}]\n" + "\n".join(rows))
    return "\n".join(parts)[:MAX_TEXT_CHARS]


def xlsx_text(content: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        parts = []
        for sheet in workbook.worksheets:
            lines = [" | ".join("" if v is None else str(v) for v in row) for row in sheet.iter_rows(values_only=True)]
            parts.append(f"[Sheet {sheet.title}]\n" + "\n".join(line.rstrip(" |") for line in lines))
        return "\n".join(parts)[:MAX_TEXT_CHARS]
    finally:
        workbook.close()


def csv_text(content: bytes) -> str:
    text = content.decode("utf-8-sig")
    try:  # tidy to one row per line with ' | ', which reads better than ragged commas
        rows = list(csv.reader(io.StringIO(text), dialect=csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")))
        return "\n".join(" | ".join(r) for r in rows)[:MAX_TEXT_CHARS]
    except csv.Error:
        return text[:MAX_TEXT_CHARS]


def prepare_image(content: bytes) -> Tuple[bytes, str]:
    """(bytes, 'jpeg' or 'png') that Bedrock accepts: upright, at most MAX_IMAGE_SIDE px, under MAX_IMAGE_BYTES.

    Phone photos carry their rotation in EXIF; the model does not read it, so the pixels are turned upright.
    """
    from PIL import Image, ImageOps

    with Image.open(io.BytesIO(content)) as im:
        keep_png = im.format == "PNG"
        im = ImageOps.exif_transpose(im)
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            background = Image.new("RGB", im.size, "white")
            background.paste(im, mask=im.getchannel("A"))
            im = background
        elif im.mode != "RGB":
            im = im.convert("RGB")
        if max(im.size) > MAX_IMAGE_SIDE:
            im.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE), Image.LANCZOS)
        if keep_png:
            buf = io.BytesIO()
            im.save(buf, "PNG", optimize=True)
            if buf.tell() <= MAX_IMAGE_BYTES:
                return buf.getvalue(), "png"
        for quality in (90, 80, 70, 60, 50):
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=quality)
            if buf.tell() <= MAX_IMAGE_BYTES:
                return buf.getvalue(), "jpeg"
        im.thumbnail((im.width // 2, im.height // 2), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=70)
        return buf.getvalue(), "jpeg"
