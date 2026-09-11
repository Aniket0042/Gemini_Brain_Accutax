"""Branded file generators for ReportSpec → PDF/CSV/XLSX/DOCX/PPTX."""
from __future__ import annotations

import csv
import io
import logging
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from gemini_brain.artifacts.report_spec import CANVAS_MAX_ROWS, XLSX_MAX_ROWS
from gemini_brain.artifacts.store import TTL_SECONDS, ArtifactRecord, put

logger = logging.getLogger("gemini_brain.artifacts.generate")

MIME = {
    "pdf": "application/pdf",
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

ACCENT = (14 / 255, 138 / 255, 117 / 255)


def _slug(title: str, fmt: str) -> str:
    base = re.sub(r"[^a-zA-Z0-9]+", "-", (title or "accutax-report")).strip("-").lower()
    base = (base or "accutax-report")[:48]
    return f"{base}.{fmt}"


def _header_bits(spec: Dict[str, Any]) -> Tuple[str, str, str]:
    title = spec.get("title") or "Accutax report"
    subtitle = spec.get("subtitle") or ""
    period = spec.get("period") or ""
    generated = spec.get("generated_at") or datetime.now(timezone.utc).isoformat()
    try:
        generated = datetime.fromisoformat(generated.replace("Z", "+00:00")).strftime("%d %b %Y %H:%M UTC")
    except ValueError:
        pass
    meta = " · ".join(p for p in (subtitle, period, f"Generated {generated}", "Confidential") if p)
    return title, subtitle, meta


def _primary_table(spec: Dict[str, Any], max_rows: int) -> Optional[Dict[str, Any]]:
    tables = spec.get("tables") or []
    if not tables:
        return None
    table = dict(tables[0])
    raw = table.get("raw_rows") or table.get("rows") or []
    table["rows"] = raw[:max_rows]
    table["truncated"] = table.get("truncated") or len(raw) > max_rows
    return table


def _chart_png(spec: Dict[str, Any]) -> Optional[bytes]:
    charts = spec.get("charts") or []
    if not charts:
        return None
    chart = charts[0]
    cats = chart.get("categories") or []
    series = chart.get("series") or []
    if not cats or not series:
        return None
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception as e:
        logger.warning("matplotlib unavailable for chart embed: %s", e)
        return None
    fig, ax = plt.subplots(figsize=(8, 3.6), dpi=120)
    x = list(range(len(cats)))
    ctype = (chart.get("chart_type") or "bar").lower()
    colors = ["#0E8A75", "#3B82F6", "#F43F5E", "#8B5CF6"]
    if ctype == "pie" and series:
        ax.pie(series[0].get("data") or [], labels=cats, colors=colors[: len(cats)])
        ax.set_title(chart.get("title") or "")
    else:
        for i, s in enumerate(series):
            data = s.get("data") or []
            color = colors[i % len(colors)]
            if ctype == "line":
                ax.plot(x, data, marker="o", label=s.get("name"), color=color)
            elif ctype == "area":
                ax.fill_between(x, data, alpha=0.25, color=color)
                ax.plot(x, data, label=s.get("name"), color=color)
            else:
                width = 0.8 / max(len(series), 1)
                offset = (i - (len(series) - 1) / 2) * width
                ax.bar([v + offset for v in x], data, width=width, label=s.get("name"), color=color)
        ax.set_xticks(x)
        ax.set_xticklabels(cats, rotation=30, ha="right")
        ax.set_ylabel(chart.get("y_label") or "AED")
        ax.legend(frameon=False)
        ax.set_title(chart.get("title") or "")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def render_csv(spec: Dict[str, Any]) -> bytes:
    table = _primary_table(spec, XLSX_MAX_ROWS)
    buf = io.StringIO()
    if not table:
        buf.write("metric,value\n")
        for kpi in spec.get("kpis") or []:
            buf.write(f"{kpi.get('label')},{kpi.get('formatted') or kpi.get('value')}\n")
        return ("\ufeff" + buf.getvalue()).encode("utf-8")
    columns = table.get("columns") or []
    keys = [c.get("key") for c in columns] or (list(table["rows"][0].keys()) if table["rows"] else [])
    writer = csv.writer(buf)
    writer.writerow([c.get("label") or k for c, k in zip(columns, keys)] or keys)
    for row in table["rows"]:
        writer.writerow([row.get(k, "") for k in keys])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def render_xlsx(spec: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    title, _, meta = _header_bits(spec)
    header_fill = PatternFill("solid", fgColor="0E8A75")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Accutax"
    summary["A1"].font = Font(bold=True, color="0E8A75", size=16)
    summary["A2"] = title
    summary["A2"].font = Font(bold=True, size=14)
    summary["A3"] = meta
    row_i = 5
    summary["A4"] = "Metric"
    summary["B4"] = "Value"
    for cell in (summary["A4"], summary["B4"]):
        cell.fill = header_fill
        cell.font = header_font
    for kpi in spec.get("kpis") or []:
        summary[f"A{row_i}"] = kpi.get("label")
        summary[f"B{row_i}"] = kpi.get("formatted") if kpi.get("value") is None else kpi.get("value")
        row_i += 1
    for note in spec.get("notes") or []:
        summary[f"A{row_i}"] = note
        row_i += 1

    data_ws = wb.create_sheet("Data")
    table = _primary_table(spec, XLSX_MAX_ROWS)
    if table:
        columns = table.get("columns") or []
        keys = [c.get("key") for c in columns]
        if not keys and table.get("raw_rows"):
            keys = list(table["raw_rows"][0].keys())
            columns = [{"key": k, "label": k.replace("_", " ").title()} for k in keys]
        for col_i, col in enumerate(columns, start=1):
            cell = data_ws.cell(1, col_i, col.get("label") or col.get("key"))
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin
        for r_i, row in enumerate(table.get("raw_rows") or table.get("rows") or [], start=2):
            if r_i - 1 > XLSX_MAX_ROWS:
                break
            for c_i, col in enumerate(columns, start=1):
                cell = data_ws.cell(r_i, c_i, row.get(col.get("key")))
                cell.border = thin
                cell.alignment = Alignment(horizontal=col.get("align") or "left")
        data_ws.freeze_panes = "A2"
        for col_i in range(1, len(columns) + 1):
            data_ws.column_dimensions[get_column_letter(col_i)].width = 18

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(spec: Dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether,
    )

    title, _, meta = _header_bits(spec)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    brand = ParagraphStyle("Brand", parent=styles["Heading1"], textColor=colors.HexColor("#0E8A75"), fontSize=11, spaceAfter=2)
    h1 = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontSize=16, spaceAfter=4, textColor=colors.HexColor("#14201C"))
    muted = ParagraphStyle("Muted", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#526059"), spaceAfter=10)
    story: List[Any] = [
        Paragraph("ACCUTAX", brand),
        Paragraph(title, h1),
        Paragraph(meta, muted),
    ]
    kpis = spec.get("kpis") or []
    if kpis:
        kpi_row = [[Paragraph(f"<b>{k.get('label')}</b><br/>{k.get('formatted') or k.get('value')}", styles["Normal"]) for k in kpis[:4]]]
        kpi_table = Table(kpi_row, colWidths=[40 * mm] * len(kpi_row[0]))
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E1F1EC")),
            ("BOX", (0, 0), (-1, -1), 0.3, colors.HexColor("#0E8A75")),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C5E4DA")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 10))

    png = _chart_png(spec)
    if png:
        img = Image(io.BytesIO(png), width=170 * mm, height=76 * mm)
        story.append(img)
        story.append(Spacer(1, 8))

    table = _primary_table(spec, CANVAS_MAX_ROWS)
    if table:
        columns = table.get("columns") or []
        keys = [c.get("key") for c in columns]
        header = [c.get("label") or k for c, k in zip(columns, keys)]
        data_rows = [header]
        for row in table.get("rows") or []:
            data_rows.append([str(row.get(k, "") or "") for k in keys])
        if len(data_rows) > 1:
            tbl = Table(data_rows, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E8A75")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D0D5DD")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(KeepTogether([tbl]))
    for note in spec.get("notes") or []:
        story.append(Paragraph(note, muted))

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#8A968E"))
        canvas.drawString(18 * mm, 10 * mm, "Accutax · Confidential")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def render_docx(spec: Dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    title, _, meta = _header_bits(spec)
    doc = Document()
    brand = doc.add_paragraph()
    run = brand.add_run("ACCUTAX")
    run.bold = True
    run.font.color.rgb = RGBColor(0x0E, 0x8A, 0x75)
    run.font.size = Pt(12)
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(meta)
    p.runs[0].font.size = Pt(9) if p.runs else None

    kpis = spec.get("kpis") or []
    if kpis:
        table = doc.add_table(rows=1, cols=min(4, len(kpis)))
        for i, kpi in enumerate(kpis[:4]):
            cell = table.rows[0].cells[i]
            cell.text = f"{kpi.get('label')}\n{kpi.get('formatted') or kpi.get('value')}"

    png = _chart_png(spec)
    if png:
        doc.add_picture(io.BytesIO(png), width=Inches(6.2))

    data_table = _primary_table(spec, CANVAS_MAX_ROWS)
    if data_table and data_table.get("columns"):
        cols = data_table["columns"]
        tbl = doc.add_table(rows=1 + len(data_table.get("rows") or []), cols=len(cols))
        tbl.style = "Table Grid"
        for i, col in enumerate(cols):
            tbl.rows[0].cells[i].text = col.get("label") or col.get("key")
        for r_i, row in enumerate(data_table.get("rows") or [], start=1):
            for c_i, col in enumerate(cols):
                tbl.rows[r_i].cells[c_i].text = str(row.get(col.get("key"), "") or "")
    for note in spec.get("notes") or []:
        doc.add_paragraph(note)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def render_pptx(spec: Dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    title, subtitle, meta = _header_bits(spec)
    prs = Presentation()
    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "\n".join(p for p in (subtitle, meta) if p)

    kpis = spec.get("kpis") or []
    if kpis:
        kpi_slide = prs.slides.add_slide(prs.slide_layouts[5])
        kpi_slide.shapes.title.text = "Key figures"
        table_shape = kpi_slide.shapes.add_table(len(kpis), 2, Inches(0.8), Inches(1.6), Inches(8.4), Inches(0.4 * max(len(kpis), 1)))
        tbl = table_shape.table
        for i, kpi in enumerate(kpis[:8]):
            tbl.cell(i, 0).text = str(kpi.get("label") or "")
            tbl.cell(i, 1).text = str(kpi.get("formatted") or kpi.get("value") or "")

    charts = spec.get("charts") or []
    if charts:
        chart = charts[0]
        cats = chart.get("categories") or []
        series = chart.get("series") or []
        if cats and series:
            ch_slide = prs.slides.add_slide(prs.slide_layouts[5])
            ch_slide.shapes.title.text = chart.get("title") or "Chart"
            data = CategoryChartData()
            data.categories = cats
            for s in series:
                data.add_series(s.get("name") or "Value", tuple(s.get("data") or []))
            ctype = (chart.get("chart_type") or "bar").lower()
            xl = XL_CHART_TYPE.COLUMN_CLUSTERED
            if ctype == "line":
                xl = XL_CHART_TYPE.LINE_MARKERS
            elif ctype == "area":
                xl = XL_CHART_TYPE.AREA
            elif ctype == "pie":
                xl = XL_CHART_TYPE.PIE
            ch_slide.shapes.add_chart(xl, Inches(0.6), Inches(1.5), Inches(8.8), Inches(5.0), data)

    table = _primary_table(spec, min(12, CANVAS_MAX_ROWS))
    if table and table.get("columns"):
        cols = table["columns"][:6]
        rows = (table.get("rows") or [])[:12]
        t_slide = prs.slides.add_slide(prs.slide_layouts[5])
        t_slide.shapes.title.text = table.get("title") or "Details"
        shape = t_slide.shapes.add_table(1 + len(rows), len(cols), Inches(0.5), Inches(1.5), Inches(9.0), Inches(0.35 * (1 + len(rows))))
        tbl = shape.table
        for i, col in enumerate(cols):
            tbl.cell(0, i).text = col.get("label") or col.get("key")
        for r_i, row in enumerate(rows, start=1):
            for c_i, col in enumerate(cols):
                tbl.cell(r_i, c_i).text = str(row.get(col.get("key"), "") or "")

    # Cap extra slides — we already created at most title + kpi + chart + table.
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


RENDERERS = {
    "csv": render_csv,
    "xlsx": render_xlsx,
    "pdf": render_pdf,
    "docx": render_docx,
    "pptx": render_pptx,
}


def generate_and_store(
    spec: Dict[str, Any],
    fmt: str,
    *,
    user_id: int,
    organization_id: Optional[int] = None,
    session_id: Optional[str] = None,
) -> Optional[Dict[str, Any]]:
    fmt = (fmt or "").lower()
    renderer = RENDERERS.get(fmt)
    if renderer is None:
        return None
    content = renderer(spec)
    if not content:
        return None
    rec: ArtifactRecord = put(
        content,
        filename=_slug(spec.get("title") or "report", fmt),
        mime=MIME[fmt],
        user_id=user_id,
        organization_id=organization_id,
        session_id=session_id,
        ttl=TTL_SECONDS,
    )
    return {
        "type": "artifact",
        "id": rec.id,
        "kind": fmt,
        "filename": rec.filename,
        "mime": rec.mime,
        "expires_in": TTL_SECONDS,
        "expires_at": datetime.fromtimestamp(rec.expires_at, tz=timezone.utc).isoformat(),
    }
