"""
golden_sheet.py — Manual test sheet for the golden accounting set, with today's reference figures.

Reads tests/data/golden_accounting.json, runs every reference Cube query for the
day of the run, and writes an Excel workbook a tester works through in the chat UI:
the question to type, the organizations to select, what a correct answer must
show, the reference figures, and Result / Notes columns to fill in.

Figures change as data is posted, so build the sheet on the day of testing.

Usage (from backend/, on a host that reaches Cube and the database):
  .venv/bin/python scripts/eval/golden_sheet.py --out /tmp/golden_manual_test.xlsx
"""
from __future__ import annotations

import argparse
import datetime
import json
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List

_BACKEND_ROOT = Path(__file__).resolve().parents[2]
for path in (_BACKEND_ROOT / "src", _BACKEND_ROOT, Path(__file__).resolve().parent):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from gemini_brain.config.settings import settings  # noqa: E402
from gemini_brain.semantic import periods  # noqa: E402
from gemini_brain.db.connection import get_connection  # noqa: E402
from replay_real_questions import reference  # noqa: E402

CASES_FILE = _BACKEND_ROOT / "tests" / "data" / "golden_accounting.json"
MAX_ROWS_SHOWN = 6
BEHAVIOR = {
    "answer": "Answer with figures",
    "partial": "Answer what the data holds; say what is missing",
    "not_available": "Say it is not available; offer the closest figure",
    "action": "Decline (read-only); point to the app",
    "clarify": "Ask one short question",
}
HEADERS = ["#", "ID", "Priority", "Team", "Select organizations", "Question (type exactly)", "Correct behaviour",
           "A correct answer must show", "Reference figures (Cube)", "Result", "Notes"]
WIDTHS = [5, 13, 8, 11, 26, 46, 26, 48, 70, 10, 30]


def _num(value: Any) -> str:
    try:
        d = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return str(value)
    return f"{d:,.2f}" if d != d.to_integral_value() else f"{d:,.0f}"


def _short(name: Any) -> str:
    text = str(name or "")
    return text.split("_User1_")[-1] if "_User1_" in text else text


def describe(ref: Dict[str, Any]) -> str:
    """A few lines a tester can compare an answer against."""
    spec = ref["spec"]
    when = spec.get("period") or (f"as of {spec['as_of']}" if spec.get("as_of") else "today")
    head = f"{spec['view']}: {', '.join(spec['measures'])} [{when}]"
    if spec.get("filters"):
        head += " where " + "; ".join(f"{m} {op} {','.join(map(str, v))}" for m, op, v in spec["filters"])
    if ref.get("error"):
        return head + "\n  (reference query failed: " + ref["error"][:120] + ")"
    rows = [r for r in ref.get("rows") or [] if any(r.get(m) is not None for m in spec["measures"])]
    if not rows:
        return head + "\n  (no rows: the answer should say nothing is recorded)"
    lines = [head]
    for row in rows[:MAX_ROWS_SHOWN]:
        label = [_short(row.get("organization_name"))] + [str(row.get(g)) for g in spec.get("group_by", [])]
        values = ", ".join(f"{m}={_num(row.get(m))}" for m in spec["measures"] if row.get(m) is not None)
        lines.append("  " + " / ".join(label) + ": " + (values or "none"))
    if len(rows) > MAX_ROWS_SHOWN:
        lines.append(f"  … {len(rows) - MAX_ROWS_SHOWN} more rows")
    additive = [m for m in spec["measures"] if not m.endswith(("_pct", "_ratio"))]
    if len(rows) > 1 and additive and len({r.get("currency") for r in rows}) == 1:
        totals = ", ".join(f"{m}={_num(sum(Decimal(str(r.get(m) or 0)) for r in rows))}" for m in additive)
        lines.append(f"  Total ({rows[0].get('currency')}): {totals}")
    return "\n".join(lines)


def write_sheet(ws, cases: List[Dict[str, Any]], names: Dict[int, str]) -> None:
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, c in enumerate(cases, 1):
        orgs = "All 10 test organizations" if len(c["orgs"]) == 10 else "\n".join(names.get(o, str(o)) for o in c["orgs"])
        ws.append([i, c["id"], c["priority"], c["team"], orgs, c["question"], BEHAVIOR[c["behavior"]],
                   c["must_show"], c.get("_reference", "—"), "", ""])
    for idx, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"Pass,Partial,Fail"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{len(cases) + 1}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--cases", default=str(CASES_FILE))
    parser.add_argument("--out", default=str(_BACKEND_ROOT / "golden_manual_test.xlsx"))
    args = parser.parse_args()

    spec = json.loads(Path(args.cases).read_text(encoding="utf-8"))
    cases = spec["cases"]
    today = periods.today_in(settings.report_timezone)

    conn = get_connection()
    cur = conn.cursor()
    cur.execute("SELECT id, name FROM organizations WHERE id = ANY(%s)", (sorted({o for c in cases for o in c["orgs"]}),))
    names = {int(i): f"{n} (id {i})" for i, n in cur.fetchall()}
    conn.close()

    for c in cases:
        refs = [reference(s, c["orgs"], today) for s in c["expect"].get("ref", [])]
        c["_reference"] = "\n\n".join(describe(r) for r in refs) if refs else "—"
        print(f"{c['id']:12} {len(refs)} reference(s)", flush=True)

    wb = Workbook()
    intro = wb.active
    intro.title = "How to test"
    for line in [
        ["Golden evaluation set — manual test sheet"],
        [f"Reference figures computed on {today:%d %b %Y} from Cube (posted ledger for P&L, documents for sales/purchases/AR/AP)."],
        ["Rebuild this sheet on the day you test: scripts/eval/golden_sheet.py"],
        [],
        ["How to run a case"],
        ["1. In the chat UI, select exactly the organizations in 'Select organizations'. Start a new chat for each case."],
        ["2. Type the question exactly as written."],
        ["3. Compare the answer with 'Correct behaviour' and 'A correct answer must show'. For figures, use 'Reference figures'."],
        ["4. Set Result: Pass (all of it), Partial (right data, something missing or unclear), Fail (wrong figure, wrong question, invented, or no answer)."],
        ["5. Note anything odd in Notes; a FireShot PDF of the answer helps the review."],
        [],
        ["Correct behaviours"],
        *[[f"{k}: {v}"] for k, v in BEHAVIOR.items()],
        [],
        ["Known data limits (not answer errors)"],
        ["Orgs 26–33 post invoices and bills to random ledger accounts, so their ledger revenue is negative. P&L cases use orgs 24–25."],
        ["Open bills in the test orgs have no due date, so 'due this week' and 'overdue payables' have nothing to show; "
         "a correct answer says the due dates are missing. No open invoice or bill is partially paid."],
        ["Org 25's ledger has postings only up to June 2026; October has no postings yet for orgs 24–25."],
        ["Draft invoices, payments/receipts, bank reconciliation, approvals and attachments are not in the governed data yet."],
        [],
        ["Sheets: 'P1 quick set' = 26 core cases (about 1 hour). 'All cases' = full set."],
    ]:
        intro.append(line)
    intro["A1"].font = Font(bold=True, size=14)
    for r in (5, 12, 19):
        intro[f"A{r}"].font = Font(bold=True)
    intro.column_dimensions["A"].width = 140

    write_sheet(wb.create_sheet("P1 quick set"), [c for c in cases if c["priority"] == "P1"], names)
    write_sheet(wb.create_sheet("All cases"), cases, names)

    out = Path(args.out)
    wb.save(out)
    print(f"Sheet: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
