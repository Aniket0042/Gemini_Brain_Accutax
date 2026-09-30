"""Phase 2: chart forms, the profit bridge, takeaways, and their rendering."""
import io
import random
from decimal import Decimal

import pytest
from pydantic import ValidationError

from gemini_brain.artifacts import theme
from gemini_brain.artifacts.attach import attach_delivery
from gemini_brain.artifacts.delivery import detect_delivery
from gemini_brain.artifacts.facts import build_factsheet
from gemini_brain.artifacts.generate import _chart_png, render_pptx, render_xlsx
from gemini_brain.artifacts.insights import add_takeaways
from gemini_brain.artifacts.integrity import ground_text
from gemini_brain.artifacts.ir import ChartSection, ReportDocument, Series, format_compact
from gemini_brain.artifacts.report_spec import build_report_spec

MONTHLY = [
    {"month": "April", "revenue": 52400, "expenses": 34100},
    {"month": "May", "revenue": 58200, "expenses": 36800},
    {"month": "June", "revenue": 63700, "expenses": 43760},
]
PNL = {
    "statement": "profit_and_loss",
    "revenue": {"total": 174300, "line_items": [{"name": "Services", "amount": 174300}]},
    "expenses": {"total": 114660, "line_items": [
        {"name": "Cost of Goods", "amount": 52290},
        {"name": "Payroll", "amount": 42600},
        {"name": "Overheads", "amount": 17430},
        {"name": "Marketing", "amount": 2340},
    ]},
    "net_profit": 59640,
    "monthly": MONTHLY,
}
CUSTOMERS = [{"customer_name": n, "sales": v} for n, v in [
    ("Apex Retail Trading LLC", 220500), ("Falcon Energy", 191310), ("Al Habtoor Group", 181755),
    ("Marina Foods", 96400), ("Desert Labs", 72210), ("Gulf Motors", 51000), ("Nakheel Co", 33000),
]]


def _spec(query, data):
    return next(b for b in attach_delivery(query, data, []) if b["type"] == "canvas")["spec"]


# ── requests ─────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("query,hint", [
    ("show the P&L waterfall", "waterfall"),
    ("profit bridge for Q2 as pdf", "waterfall"),
    ("stacked bar of vat and net sales by month", "stacked_bar"),
    ("horizontal bar of sales by customer", "hbar"),
    ("P&L report as pdf", None),  # no named form: the data decides
])
def test_chart_form_requests(query, hint):
    assert detect_delivery(query).chart_hint == hint


def test_a_company_called_bridge_is_not_a_chart_request():
    assert detect_delivery("Bridge Corp sales").mode == "none"


# ── choosing the form from the data ──────────────────────────────────────────

def test_ranking_of_many_or_long_names_is_horizontal():
    assert _spec("sales by customer chart", CUSTOMERS)["charts"][0]["chart_type"] == "hbar"


def test_short_ranking_stays_vertical():
    rows = [{"name": "A", "amount": 3}, {"name": "B", "amount": 2}, {"name": "C", "amount": 1}]
    assert _spec("chart of amount by name", rows)["charts"][0]["chart_type"] == "bar"


def test_share_words_give_a_donut():
    assert _spec("breakdown chart of sales by customer", CUSTOMERS)["charts"][0]["chart_type"] == "donut"


def test_an_explicit_bar_request_is_never_overridden():
    spec = _spec("bar chart of sales by customer", CUSTOMERS)
    assert spec["charts"][0]["chart_type"] == "bar"


def test_time_series_keeps_its_form():
    spec = _spec("revenue trend line", [{"month": m["month"], "revenue": m["revenue"]} for m in MONTHLY])
    assert spec["charts"][0]["chart_type"] == "line"


# ── the profit bridge ────────────────────────────────────────────────────────

def test_pnl_gets_a_bridge_that_adds_up():
    spec = _spec("P&L report as pdf", PNL)
    bridge = next(c for c in spec["charts"] if c["chart_type"] == "waterfall")
    assert bridge["categories"] == ["Revenue", "Cost of Goods", "Payroll", "Overheads", "Marketing", "Net Profit"]
    assert bridge["series"][0]["data"] == [174300.0, -52290.0, -42600.0, -17430.0, -2340.0, 59640.0]
    assert bridge["total_indices"] == [0, 5]
    assert bridge["point_colors"][0] == theme.PRIMARY and bridge["point_colors"][1] == theme.CATEGORICAL[2]


def test_waterfall_request_puts_the_bridge_first():
    assert _spec("P&L waterfall", PNL)["charts"][0]["chart_type"] == "waterfall"


def test_unexplained_gap_is_its_own_step_with_a_note():
    data = dict(PNL, net_profit=55000)
    spec = _spec("P&L waterfall", data)
    bridge = spec["charts"][0]
    assert "Other income / costs" in bridge["categories"]
    assert bridge["series"][0]["data"][-2] == -4640.0
    assert any("do not fully explain the gap" in n for n in spec["notes"])


def test_many_costs_fold_into_other_costs():
    items = [{"name": f"Cost {i}", "amount": 1000 * (10 - i)} for i in range(9)]
    data = {"statement": "profit_and_loss", "revenue": {"total": 100000, "line_items": [{"name": "S", "amount": 100000}]},
            "expenses": {"total": sum(i["amount"] for i in items), "line_items": items}, "net_profit": 100000 - 54000}
    bridge = _spec("P&L waterfall", data)["charts"][0]
    assert bridge["categories"][-2] == "Other costs" and len(bridge["categories"]) == 8


def test_waterfall_without_bridge_data_falls_back_with_reason():
    spec = _spec("waterfall chart of sales by customer", CUSTOMERS)
    assert spec["charts"][0]["chart_type"] == "bar"
    assert any("waterfall needs a starting total" in n for n in spec["notes"])


def test_bridge_validator_rejects_steps_that_do_not_add_up():
    with pytest.raises(ValidationError, match="does not add up"):
        ChartSection(chart_type="waterfall", categories=["Revenue", "Costs", "Net Profit"],
                     series=[Series(name="Amount", values=[Decimal(100), Decimal(-30), Decimal(80)])],
                     total_indices=[0, 2])


def test_total_bars_only_on_waterfalls():
    with pytest.raises(ValidationError, match="only a waterfall"):
        ChartSection(chart_type="bar", categories=["A", "B"],
                     series=[Series(name="Sales", values=[Decimal(1), Decimal(2)])], total_indices=[0])


# ── stacked ──────────────────────────────────────────────────────────────────

def test_stacked_parts():
    rows = [{"month": "April", "vat": 100, "net_sales": 2000}, {"month": "May", "vat": 120, "net_sales": 2400}]
    chart = _spec("stacked bar of vat and net sales by month", rows)["charts"][0]
    assert chart["chart_type"] == "stacked_bar"
    assert [s["name"] for s in chart["series"]] == ["Net Sales", "VAT"]


def test_revenue_and_expenses_are_never_stacked():
    spec = _spec("stacked bar of revenue and expenses by month", {"statement": "profit_and_loss", "monthly": MONTHLY})
    assert spec["charts"][0]["chart_type"] == "bar"
    assert any("shown side by side" in n for n in spec["notes"])


def test_stacked_rejects_negative_parts():
    with pytest.raises(ValidationError, match="negative"):
        ChartSection(chart_type="stacked_bar", categories=["A"],
                     series=[Series(name="X", values=[Decimal(1)]), Series(name="Y", values=[Decimal(-1)])])


# ── takeaways ────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("query,data,index,expected", [
    ("sales by customer chart", CUSTOMERS, 0,
     "Apex Retail Trading LLC is the largest at AED 220.5K, 26.1% of the total. The top three make up 70.1%."),
    ("revenue trend line", [{"month": m["month"], "revenue": m["revenue"]} for m in MONTHLY], 0,
     "Revenue rose 21.6% from April to June, from AED 52.4K to AED 63.7K."),
    ("P&L report as pdf", PNL, 0,
     "Revenue was above Expenses in 3 of 3 periods; the widest gap was AED 21.4K in May."),
    ("P&L report as pdf", PNL, 1,
     "Costs of AED 114.7K take revenue of AED 174.3K to net profit of AED 59.6K; "
     "Cost of Goods is the largest cost at AED 52.3K."),
])
def test_takeaways_say_the_right_thing(query, data, index, expected):
    assert _spec(query, data)["charts"][index]["takeaway"] == expected


def test_two_figures_compare_instead_of_sharing():
    spec = _spec("revenue vs expenses chart", {"statement": "profit_and_loss", "total_revenue": 1000, "total_expenses": 550})
    assert spec["charts"][0]["takeaway"] == "Revenue is AED 450 above Expenses (AED 1,000 against AED 550)."


def test_unrelated_kpis_get_no_takeaway():
    spec = _spec("chart of totals", {"total_income": 5000, "total_expense": 3000, "tax_due": 400})
    assert all(not c.get("takeaway") for c in spec["charts"] if c["title"] == "Key figures")


def test_every_takeaway_grounds_against_its_data():
    """Takeaways are computed, so they must always pass the grounding check."""
    rng = random.Random(11)
    for _ in range(30):
        n = rng.randint(2, 12)
        rows = [{"customer_name": f"Client {chr(65 + i)}", "sales": rng.randint(500, 2_000_000)} for i in range(n)]
        months = [{"month": m, "revenue": rng.randint(1_000, 500_000), "expenses": rng.randint(1_000, 500_000)}
                  for m in ("January", "February", "March", "April")[: rng.randint(2, 4)]]
        for query, data, hint in (
            ("sales by customer", rows, None),
            ("pie of sales by customer", rows, "pie"),
            ("revenue by month", [{"month": m["month"], "revenue": m["revenue"]} for m in months], "line"),
            ("P&L", {"statement": "profit_and_loss", "monthly": months}, None),
        ):
            doc = add_takeaways(ReportDocument.from_spec(build_report_spec(data, query, chart_hint=hint)))
            facts = build_factsheet(doc)
            for chart in doc.charts:
                if chart.takeaway:
                    # Computed takeaways are grounded unscoped, as verify_document does.
                    assert ground_text(chart.takeaway, facts, scoped=False).removed == [], (chart.takeaway, data)


# ── rendering ────────────────────────────────────────────────────────────────

def test_compact_format():
    assert format_compact(Decimal("220500"), "money", "AED") == "AED 220.5K"
    assert format_compact(Decimal("1240000"), "money", "AED") == "AED 1.24M"
    assert format_compact(Decimal("2340"), "money", "AED") == "AED 2,340"
    assert format_compact(Decimal("-52290"), "number") == "-52.3K"


@pytest.mark.parametrize("query,data", [
    ("P&L report as pdf", PNL),
    ("sales by customer chart", CUSTOMERS),
    ("donut of sales by customer", CUSTOMERS),
    ("stacked bar of vat and net sales by month",
     [{"month": "April", "vat": 100, "net_sales": 2000}, {"month": "May", "vat": 120, "net_sales": 2400}]),
    ("revenue trend line", [{"month": m["month"], "revenue": m["revenue"]} for m in MONTHLY]),
])
def test_every_chart_type_renders_to_png(query, data):
    for chart in _spec(query, data)["charts"]:
        png = _chart_png(chart)
        assert png and png[:4] == b"\x89PNG", chart["chart_type"]


def test_native_office_charts_use_the_same_forms():
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE

    pnl = _spec("P&L report as pdf", PNL)
    prs = Presentation(io.BytesIO(render_pptx(pnl)))
    types = [s.chart.chart_type for slide in prs.slides for s in slide.shapes if s.has_chart]
    assert XL_CHART_TYPE.COLUMN_STACKED in types  # the bridge, over an invisible base

    ranking = _spec("sales by customer chart", CUSTOMERS)
    prs = Presentation(io.BytesIO(render_pptx(ranking)))
    chart = next(s.chart for slide in prs.slides for s in slide.shapes if s.has_chart)
    assert chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
    assert chart.category_axis.reverse_order is True
    texts = [s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame]
    assert ranking["charts"][0]["takeaway"] in texts

    wb = load_workbook(io.BytesIO(render_xlsx(pnl)))
    groupings = [(c.type, c.grouping) for c in wb["Charts"]._charts]
    assert ("col", "stacked") in groupings
