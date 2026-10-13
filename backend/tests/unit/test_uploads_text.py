"""Text out of files with no model, and the image shrinking Bedrock needs."""
import io

from PIL import Image

from gemini_brain.uploads import text
from tests.unit import upload_files as f


def test_a_text_pdf_gives_its_text_page_by_page():
    pages = text.pdf_pages(f.text_pdf(f.BILL_LINES, pages=2))
    assert len(pages) == 2 and "TAX INVOICE" in pages[0] and "525.00" in pages[1]
    assert not text.is_scanned(pages)


def test_a_pdf_without_a_text_layer_is_scanned():
    assert text.is_scanned(text.pdf_pages(f.pdf(3)))
    assert text.is_scanned([])  # unreadable: treat as a picture and let the model look


def test_an_unreadable_pdf_gives_no_pages():
    assert text.pdf_pages(b"%PDF-1.4 garbage") == []


def test_bill_words():
    assert text.looks_like_bill("TAX INVOICE\nTRN: 100\nTotal 5.00")
    assert text.looks_like_bill("Statement of Account for March. Invoice 12.")
    assert not text.looks_like_bill("This agreement is made between the parties. Payment terms: thirty days.")
    assert not text.looks_like_bill("An invoice is mentioned once.")  # one word is not enough


def test_arabic_bill_words():
    assert text.looks_like_bill("فاتورة ضريبية\nضريبة القيمة المضافة")


def test_join_pages_labels_pages_and_skips_empty_ones():
    assert text.join_pages(["one", "  ", "three"]) == "[Page 1]\none\n\n[Page 3]\nthree"


def test_docx_text_has_paragraphs_and_tables():
    import docx as docx_lib

    d = docx_lib.Document()
    d.add_paragraph("Payment terms: 30 days.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Radiator", "55"
    buf = io.BytesIO()
    d.save(buf)
    out = text.docx_text(buf.getvalue())
    assert "Payment terms: 30 days." in out and "[Table 1]" in out and "Radiator | 55" in out


def test_xlsx_text_has_each_sheet():
    out = text.xlsx_text(f.xlsx(sheets=2, rows=2))
    assert "[Sheet Sheet]" in out and "[Sheet Sheet1]" in out and "1 | line 1 | 1.5" in out


def test_csv_text_reads_commas_semicolons_and_a_bom():
    assert text.csv_text("date,ref\n2026-01-01,R1\n".encode("utf-8-sig")) == "date | ref\n2026-01-01 | R1"
    assert text.csv_text(b"a;b\n1;2\n") == "a | b\n1 | 2"
    assert text.csv_text(b"just one column\nnext\n").startswith("just one column")


# ── images ──

def _jpeg_with_orientation(orientation: int, size=(200, 100)) -> bytes:
    im = Image.new("RGB", size, "white")
    exif = Image.Exif()
    exif[0x0112] = orientation
    buf = io.BytesIO()
    im.save(buf, "JPEG", exif=exif)
    return buf.getvalue()


def test_a_phone_photo_is_turned_upright():
    data, fmt = text.prepare_image(_jpeg_with_orientation(6))   # rotated 90 degrees in EXIF
    assert fmt == "jpeg" and Image.open(io.BytesIO(data)).size == (100, 200)


def test_a_large_image_is_shrunk_to_the_limit():
    data, fmt = text.prepare_image(f.image("PNG", size=(7000, 5000)))
    assert max(Image.open(io.BytesIO(data)).size) == text.MAX_IMAGE_SIDE and len(data) <= text.MAX_IMAGE_BYTES


def test_a_transparent_png_gets_a_white_background():
    im = Image.new("RGBA", (50, 50), (0, 0, 0, 0))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    data, fmt = text.prepare_image(buf.getvalue())
    assert Image.open(io.BytesIO(data)).convert("RGB").getpixel((5, 5)) == (255, 255, 255)


def test_a_noisy_image_is_compressed_under_the_byte_limit():
    import numpy as np

    noise = Image.fromarray(np.random.default_rng(1).integers(0, 255, (2500, 2500, 3), dtype=np.uint8))
    buf = io.BytesIO()
    noise.save(buf, "PNG")
    data, fmt = text.prepare_image(buf.getvalue())
    assert fmt == "jpeg" and len(data) <= text.MAX_IMAGE_BYTES
