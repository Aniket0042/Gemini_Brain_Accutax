"""Phase 0 report fixes: requested chart type honored, one color system,
missing is not zero, tenant-scoped row links, durable artifact links."""
import io
import re
from pathlib import Path

import pytest

from gemini_brain.artifacts import store, theme
from gemini_brain.artifacts.attach import attach_delivery
from gemini_brain.artifacts.generate import _chart_png, render_pptx, render_xlsx
from gemini_brain.artifacts.report_spec import _as_float, _section_total, build_report_spec
from gemini_brain.tools import formatters

CUSTOMERS = [
    {"customer_name": f"Customer {i}", "sales": 1000.0 * (13 - i), "tax": 50.0 * (13 - i), "invoice_count": i}
    for i in range(1, 13)
]


# ── Requested chart type ─────────────────────────────────────────────────────

def test_pie_request_with_several_numeric_columns_is_a_pie():
    """Rows with sales + tax used to make 2 series, which silently forced bar."""
    rows = CUSTOMERS[:4]
    spec = build_report_spec(rows, "pie chart of sales by customer", chart_hint="pie")
    chart = spec["charts"][0]
    assert chart["chart_type"] == "pie"
    assert len(chart["series"]) == 1
    assert chart["series"][0]["name"] == "Sales"


def test_pie_uses_the_measure_the_user_named():
    spec = build_report_spec(CUSTOMERS[:4], "pie chart of tax by customer", chart_hint="pie")
    chart = spec["charts"][0]
    assert chart["chart_type"] == "pie"
    assert chart["series"][0]["name"] == "Tax"


def test_donut_request_is_a_donut():
    spec = build_report_spec(CUSTOMERS[:4], "donut of sales by customer", chart_hint="donut")
    assert spec["charts"][0]["chart_type"] == "donut"


def test_pie_with_many_categories_folds_tail_into_other_and_says_so():
    spec = build_report_spec(CUSTOMERS, "pie chart of top 12 customers by sales", chart_hint="pie")
    chart = spec["charts"][0]
    assert chart["chart_type"] == "pie"
    assert len(chart["categories"]) == 8
    assert chart["categories"][-1] == "Other"
    assert sum(chart["series"][0]["data"]) == pytest.approx(sum(r["sales"] for r in CUSTOMERS))
    assert any("grouped into Other" in n for n in spec["notes"])


def test_pie_with_negative_values_falls_back_to_bar_with_a_reason():
    rows = [{"account_name": "A", "net_amount": 500}, {"account_name": "B", "net_amount": -200}]
    spec = build_report_spec(rows, "pie of net amount by account", chart_hint="pie")
    chart = spec["charts"][0]
    assert chart["chart_type"] == "bar"
    assert chart["requested_type"] == "pie"
    assert "negative" in chart["downgrade_reason"]
    assert any("shown as a bar chart" in n for n in spec["notes"])


def test_pie_of_monthly_income_and_expense_slices_period_totals():
    data = {"graphData": {
        "labels": ["Jan 2026", "Feb 2026"],
        "incomeValues": [100, 300],
        "expenseValues": [50, 50],
        "cashflowValues": [50, 250],
    }}
    spec = build_report_spec(data, "pie chart of income vs expense", chart_hint="pie")
    chart = spec["charts"][0]
    assert chart["chart_type"] == "pie"
    assert chart["categories"] == ["Income", "Expense"]
    assert chart["series"][0]["data"] == [400.0, 100.0]
    assert any("Cash Flow" in n for n in spec["notes"])


def test_chat_card_chart_matches_requested_type():
    """A formatter line chart already in the chat blocks must not win over the pie request."""
    formatter_chart = {
        "type": "chart", "chart_type": "line", "x_label": "Month",
        "categories": ["Jan 2026", "Feb 2026"],
        "series": [{"name": "Income", "data": [100, 300]}, {"name": "Expense", "data": [50, 50]}],
    }
    out = attach_delivery("show income vs expense as a pie chart", {"total_income": 400}, [formatter_chart])
    chat_charts = [b for b in out if b["type"] == "chart"]
    canvas = next(b for b in out if b["type"] == "canvas")["spec"]
    assert chat_charts[0]["chart_type"] == "pie"
    assert canvas["charts"][0]["chart_type"] == "pie"
    assert chat_charts[0]["point_colors"] == canvas["charts"][0]["point_colors"]


# ── One color system ─────────────────────────────────────────────────────────

def test_semantic_series_colors_are_stable():
    chart = theme.resolve_chart_colors({
        "chart_type": "bar", "categories": ["Jan", "Feb"],
        "series": [{"name": "Expenses", "data": [1, 2]}, {"name": "Revenue", "data": [3, 4]}],
    })
    assert chart["series_colors"] == [theme.CATEGORICAL[2], theme.CATEGORICAL[0]]


def test_single_series_bar_is_one_color():
    rows = [{"customer_name": r["customer_name"], "sales": r["sales"]} for r in CUSTOMERS[:5]]
    spec = build_report_spec(rows, "bar chart of sales by customer", chart_hint="bar")
    chart = spec["charts"][0]
    assert chart["series_colors"] == [theme.PRIMARY]
    assert "point_colors" not in chart


def test_revenue_vs_expense_bars_colored_by_meaning():
    chart = theme.resolve_chart_colors({
        "chart_type": "bar", "categories": ["Revenue", "Expenses"],
        "series": [{"name": "Amount", "data": [10, 4]}],
    })
    assert chart["point_colors"] == [theme.CATEGORICAL[0], theme.CATEGORICAL[2]]


def test_pie_slices_never_repeat_a_color():
    spec = build_report_spec(CUSTOMERS, "pie of sales by customer", chart_hint="pie")
    colors = spec["charts"][0]["point_colors"]
    assert len(set(colors)) == len(colors)
    assert colors[-1] == theme.OTHER_COLOR


def test_pptx_and_xlsx_use_the_spec_colors():
    from openpyxl import load_workbook
    from pptx import Presentation

    spec = build_report_spec(CUSTOMERS[:4], "pie of sales by customer", chart_hint="pie")
    expected = [theme.bare(c).upper() for c in spec["charts"][0]["point_colors"]]

    prs = Presentation(io.BytesIO(render_pptx(spec)))
    chart = next(s.chart for slide in prs.slides for s in slide.shapes if s.has_chart)
    got = [str(chart.plots[0].series[0].points[i].format.fill.fore_color.rgb) for i in range(len(expected))]
    assert got == expected

    wb = load_workbook(io.BytesIO(render_xlsx(spec)))
    xl_chart = wb["Charts"]._charts[0]
    fills = [dp.graphicalProperties.solidFill for dp in xl_chart.series[0].dPt]
    got_xl = [str(getattr(f, "srgbClr", f) if not isinstance(f, str) else f).upper() for f in fills]
    assert got_xl == expected


def test_frontend_palette_mirror_matches_backend():
    js = Path(__file__).resolve().parents[3] / "frontend" / "src" / "components" / "canvas" / "chartTheme.js"
    if not js.exists():
        pytest.skip("frontend not checked out")
    text = js.read_text(encoding="utf-8")
    block = re.search(r"CATEGORICAL\s*=\s*\[(.*?)\]", text, re.S).group(1)
    assert [h.upper() for h in re.findall(r"#[0-9A-Fa-f]{6}", block)] == [c.upper() for c in theme.CATEGORICAL]
    assert theme.OTHER_COLOR.upper() in text.upper()


# ── Missing is not zero; numbers parse correctly ─────────────────────────────

def test_zero_revenue_is_not_replaced_by_gross_profit():
    spec = build_report_spec(
        {"statement": "profit and loss", "total_revenue": 0, "gross_profit": 900, "total_expenses": 100},
        "P&L",
    )
    kpis = {k["label"]: k["value"] for k in spec["kpis"]}
    assert kpis["Total Revenue"] == 0
    assert kpis["Net Profit"] == -100


def test_operating_profit_keeps_its_own_label():
    spec = build_report_spec({"statement": "profit and loss", "operatingProfit": 700}, "P&L")
    labels = {k["label"] for k in spec["kpis"]}
    assert "Operating Profit" in labels
    assert "Net Profit" not in labels


def test_missing_month_is_a_gap_not_zero():
    data = {"statement": "profit and loss", "monthly": [
        {"month": "Jan 2026", "revenue": 100, "expenses": 40},
        {"month": "Feb 2026", "revenue": None, "expenses": 30},
    ]}
    spec = build_report_spec(data, "P&L")
    revenue = next(s for s in spec["charts"][0]["series"] if s["name"] == "Revenue")
    assert revenue["data"] == [100.0, None]
    monthly = next(t for t in spec["tables"] if t["title"] == "Monthly Breakdown")
    assert monthly["rows"][1]["revenue"] == "—"
    assert monthly["rows"][1]["net_profit"] == "—"


def test_accounting_negative_parses():
    assert _as_float("(1,234.50)") == -1234.5
    assert _as_float("AED (200)") == -200.0
    assert _as_float("—") is None


def test_section_total_skips_subtotal_rows():
    section = {"line_items": [
        {"name": "Consulting", "amount": 100},
        {"name": "Licences", "amount": 50},
        {"name": "Subtotal", "amount": 150},
    ]}
    assert _section_total(section) == 150


def test_chart_png_never_raises_on_bad_data():
    png = _chart_png({"chart_type": "pie", "title": "x", "categories": ["a", "b"],
                      "series": [{"name": "v", "data": [None, None]}]})
    assert png is None or png[:4] == b"\x89PNG"


# ── Tenant-scoped row links ──────────────────────────────────────────────────

def test_row_link_lookup_without_tenant_never_queries(monkeypatch):
    import gemini_brain.sql_fallback.db_connection as db

    def boom(*a, **k):
        raise AssertionError("must not connect without an organization")

    monkeypatch.setattr(db, "get_connection", boom)
    block = formatters.render_table_block([{"invoice_number": "INV-1", "amount": 10}])
    assert block["rows"] and "_row_id" not in block["rows"][0]


def test_row_link_lookup_is_filtered_by_organization(monkeypatch):
    import gemini_brain.sql_fallback.db_connection as db

    executed = []

    class Cur:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def execute(self, sql, params=None):
            executed.append((sql, params))

        def fetchall(self):
            return [("INV-1", 77)]

    class Conn:
        def cursor(self):
            return Cur()

        def rollback(self):
            pass

        def close(self):
            pass

    monkeypatch.setattr(db, "get_connection", lambda db_name="": Conn())
    blocks = formatters.render_blocks("row_table", [{"invoice_number": "INV-1", "amount": 10}], organization_id=42)
    assert blocks[0]["rows"][0]["_row_id"] == 77
    select_sql, params = executed[-1]
    assert "organization_id = %s" in select_sql
    assert params[0] == 42


# ── Durable artifact links ───────────────────────────────────────────────────

def test_artifact_survives_process_cache_loss():
    store.reset_for_tests()
    try:
        rec = store.put(b"x", filename="a.pdf", mime="application/pdf", user_id=3, organization_id=9, ttl=3600)
        store._records.clear()  # what a restart or another worker sees
        again = store.get(rec.id)
        assert again is not None and again.user_id == 3 and again.organization_id == 9
        assert store.get("../../etc/passwd") is None
    finally:
        store.reset_for_tests()


def test_default_ttl_is_long_enough_to_download():
    assert store.TTL_SECONDS >= 3600
