"""Renderer tests: every format produces valid bytes and no format drops data.

Regression target: render_pdf/docx/xlsx/pptx used to read only tables[0] and
charts[0], silently dropping the rest of a multi-table spec (e.g. a P&L's
"Monthly Breakdown" table alongside its "Line Items" table).
"""
import io

from gemini_brain.artifacts.generate import MIME, RENDERERS, _human_axis_value
from gemini_brain.artifacts.report_spec import build_report_spec

PL_PAYLOAD = {
    "statement": "profit_and_loss",
    "revenue": {
        "line_items": [
            {"name": "Product Sales", "amount": 120000},
            {"name": "Service Revenue", "amount": 45000},
        ],
        "total": 165000,
    },
    "expenses": {
        "line_items": [
            {"name": "Salaries", "amount": 60000},
            {"name": "Rent", "amount": 15000},
        ],
        "total": 75000,
    },
    "monthly": [
        {"month": "Jan", "revenue": 50000, "expenses": 20000},
        {"month": "Feb", "revenue": 60000, "expenses": 25000},
        {"month": "Mar", "revenue": 55000, "expenses": 30000},
    ],
}


def _pl_spec():
    spec = build_report_spec(PL_PAYLOAD, "profit and loss report")
    # Sanity on the fixture itself — if this stops holding, the renderer
    # assertions below stop testing the multi-table/chart case they claim to.
    assert len(spec["tables"]) == 2
    assert len(spec["charts"]) >= 1
    return spec


def _empty_spec():
    return build_report_spec({}, "")


def test_mime_covers_every_renderer():
    assert set(RENDERERS) == set(MIME)


def test_every_format_renders_nonempty_bytes_for_full_spec():
    spec = _pl_spec()
    for fmt, renderer in RENDERERS.items():
        content = renderer(spec)
        assert content, f"{fmt} renderer returned empty bytes"
        assert isinstance(content, bytes)


def test_every_format_tolerates_empty_spec():
    spec = _empty_spec()
    for fmt, renderer in RENDERERS.items():
        content = renderer(spec)
        assert content, f"{fmt} renderer returned empty bytes for an empty spec"


def test_file_signatures():
    spec = _pl_spec()
    assert RENDERERS["pdf"](spec).startswith(b"%PDF")
    for fmt in ("xlsx", "docx", "pptx"):
        assert RENDERERS[fmt](spec).startswith(b"PK\x03\x04"), fmt
    assert RENDERERS["md"](spec).decode("utf-8").lstrip().startswith("**ACCUTAX**")
    csv_text = RENDERERS["csv"](spec).decode("utf-8-sig")
    assert csv_text.splitlines()[0] == "Month,Revenue,Expenses,Net Profit"


def test_xlsx_gets_one_sheet_per_table_plus_native_chart():
    from openpyxl import load_workbook

    content = RENDERERS["xlsx"](_pl_spec())
    wb = load_workbook(io.BytesIO(content))
    assert {"Summary", "Monthly Breakdown", "Line Items", "Charts"}.issubset(set(wb.sheetnames))
    charts_ws = wb["Charts"]
    assert len(charts_ws._charts) >= 1  # native, editable chart object — not a picture


def test_docx_keeps_both_tables():
    from docx import Document

    content = RENDERERS["docx"](_pl_spec())
    doc = Document(io.BytesIO(content))
    # KPI table + Monthly Breakdown table + Line Items table
    assert len(doc.tables) >= 3


def test_pptx_gets_a_slide_per_table_and_chart():
    from pptx import Presentation

    content = RENDERERS["pptx"](_pl_spec())
    prs = Presentation(io.BytesIO(content))
    # Title + KPI + >=1 chart + 2 table slides
    assert len(prs.slides) >= 5


def test_md_embeds_both_tables_and_a_chart_image():
    text = RENDERERS["md"](_pl_spec()).decode("utf-8")
    assert "## Monthly Breakdown" in text
    assert "## Line Items" in text
    assert "data:image/png;base64," in text
    assert "| Metric | Value |" in text


def test_parse_narrative_lines_classifies_bullets_headings_paragraphs():
    from gemini_brain.artifacts.generate import _parse_narrative_lines

    text = "Revenue is up.\n\n- First point\n* Second point\n\n## Watch\nSomething to watch."
    parsed = _parse_narrative_lines(text)
    assert parsed[0] == ("para", "Revenue is up.")
    assert parsed[1] == ("bullet", "First point")
    assert parsed[2] == ("bullet", "Second point")
    assert parsed[3] == ("heading", "Watch")
    assert parsed[4] == ("para", "Something to watch.")


def test_escape_and_bold_converts_markdown_bold_to_reportlab_markup():
    from gemini_brain.artifacts.generate import _escape_and_bold

    assert _escape_and_bold("Revenue **up** 12% <ok>") == "Revenue <b>up</b> 12% &lt;ok&gt;"


def test_narrative_reaches_docx_pptx_and_md():
    """The chat's full written answer must ride along in the exported file,
    not just KPIs/charts/tables with no explanation."""
    spec = _pl_spec()
    spec["narrative"] = "Revenue climbed this quarter.\n\n- Product sales led growth\n- Watch operating costs"

    from docx import Document

    docx_doc = Document(io.BytesIO(RENDERERS["docx"](spec)))
    docx_text = "\n".join(p.text for p in docx_doc.paragraphs)
    assert "Revenue climbed this quarter." in docx_text
    assert "Product sales led growth" in docx_text

    from pptx import Presentation

    prs = Presentation(io.BytesIO(RENDERERS["pptx"](spec)))
    slide_titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert "Executive summary" in slide_titles

    md_text = RENDERERS["md"](spec).decode("utf-8")
    assert "Revenue climbed this quarter." in md_text
    assert "- Product sales led growth" in md_text


def test_human_axis_value_matches_canvas_formatting():
    """Regression: matplotlib's default scientific-notation offset ("1e6" in
    a corner) is unreadable in a static PNG. Values must read like the live
    canvas's own axis (1.2M / 800K), not raw floats or 'e' notation."""
    assert _human_axis_value(1_200_000) == "1.2M"
    assert _human_axis_value(1_000_000) == "1M"
    assert _human_axis_value(800_000) == "800K"
    assert _human_axis_value(-400_000) == "-400K"
    assert _human_axis_value(0) == "0"
    assert _human_axis_value(250) == "250"
