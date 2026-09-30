"""Phase 5: every format carries the whole report, laid out properly."""
import datetime
import io

import pytest

from gemini_brain.artifacts.attach import attach_delivery
from gemini_brain.artifacts.generate import (
    _column_widths,
    _pdf_fonts,
    render_docx,
    render_md,
    render_pdf,
    render_pptx,
    render_xlsx,
)

MONTHLY = [
    {"month": "April", "revenue": 52400, "expenses": 34100},
    {"month": "May", "revenue": 58200, "expenses": 36800},
    {"month": "June", "revenue": 63700, "expenses": 43760},
]
PNL = {
    "statement": "profit_and_loss",
    "period": {"label": "Q2 2026"},
    "revenue": {"total": 174300, "line_items": [{"name": "Services", "amount": 174300}]},
    "expenses": {"total": 114660, "line_items": [
        {"name": "Cost of Goods", "amount": 52290},
        {"name": "Payroll", "amount": 42600},
        {"name": "Overheads", "amount": 17430},
        {"name": "Marketing", "amount": 2340},
    ]},
    "net_profit": 59640,
    "margin_pct": 34.2,
    "monthly": MONTHLY,
}
INVOICES = [
    {"invoice_number": f"INV-{1000 + i}", "customer_name": f"Customer with a fairly long trading name {i} LLC",
     "issue_date": "2026-03-01", "due_date": "2026-04-01", "amount": 1234.5 * i, "status": "UNPAID"}
    for i in range(1, 80)
]


def _spec(query, data):
    return next(b for b in attach_delivery(query, data, []) if b["type"] == "canvas")["spec"]


# ── PDF ──────────────────────────────────────────────────────────────────────

def test_pdf_embeds_a_unicode_font():
    assert _pdf_fonts()[0] == "ReportSans"
    spec = _spec("P&L report as pdf", dict(PNL, entity="شركة النخيل ₹ Trading"))
    assert b"DejaVu" in render_pdf(spec)


def test_pdf_pages_are_numbered_and_long_tables_say_they_are_cut():
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open(stream=render_pdf(_spec("overdue invoices as pdf", INVOICES)))
    text = "\n".join(page.get_text() for page in doc)
    assert f"Page 1 of {len(doc)}" in text
    assert "Showing 50 of 79 rows" in text
    assert "Overdue Invoices" in doc[1].get_text()  # running header on later pages


def test_short_columns_never_shrink():
    from reportlab.pdfbase.pdfmetrics import stringWidth

    rows = [["Invoice Number", "Status", "Customer Name", "Amount"]] + [
        [f"INV-{i}", "UNPAID", "A customer with a very long registered trading name LLC", "AED 97,525.50"]
        for i in range(20)
    ]
    font = _pdf_fonts()[1]
    widths = _column_widths(rows, font, 8, 300)
    assert sum(widths) <= 300.01
    for col in (1, 3):  # status and amount keep their full width
        assert widths[col] >= stringWidth(rows[1][col], font, 8) + 12 - 0.01
    assert widths[2] < stringWidth(rows[1][2], font, 8)  # the long name wraps instead


# ── DOCX ─────────────────────────────────────────────────────────────────────

def test_docx_keeps_every_kpi_and_right_aligns_numbers():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    kpis = {"total_income": 5000, "total_expense": 3000, "tax_due": 400,
            "invoice_count": 12, "bill_count": 7, "customer_count": 5}
    doc = Document(io.BytesIO(render_docx(_spec("summary as word document", kpis))))
    text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for label in ("TOTAL INCOME", "TAX DUE", "BILL COUNT", "CUSTOMER COUNT"):
        assert label in text

    doc = Document(io.BytesIO(render_docx(_spec("P&L report as word document", PNL))))
    assert abs(doc.sections[0].page_width.cm - 21.0) < 0.1  # A4
    assert "Profit & Loss Statement" in doc.sections[0].header.paragraphs[0].text
    monthly = next(t for t in doc.tables if t.rows[0].cells[0].text == "Month")
    assert monthly.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


# ── PPTX ─────────────────────────────────────────────────────────────────────

def test_pptx_is_widescreen_and_nothing_leaves_the_slide():
    from pptx import Presentation

    prs = Presentation(io.BytesIO(render_pptx(_spec("P&L report as pptx", PNL))))
    assert round(prs.slide_width / prs.slide_height, 2) == 1.78
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (n, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 1, (n, shape.name)
            assert shape.top + shape.height <= prs.slide_height + 1, (n, shape.name)


def test_pptx_long_tables_continue_across_slides():
    from pptx import Presentation

    prs = Presentation(io.BytesIO(render_pptx(_spec("overdue invoices as pptx", INVOICES))))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]
    assert "Details" in titles and "Details (continued)" in titles
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert any("Showing 48 of 79 rows" in t for t in texts)


def test_pptx_summary_slide_leads_with_the_headline():
    from pptx import Presentation

    spec = _spec("P&L report as pptx", PNL)
    prs = Presentation(io.BytesIO(render_pptx(spec)))
    slide = next(s for s in prs.slides if s.shapes.title is not None and s.shapes.title.text == "Executive summary")
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert any(t.startswith(spec["narrative_parts"]["headline"]) for t in texts)


# ── XLSX ─────────────────────────────────────────────────────────────────────

def test_xlsx_stores_real_values_with_live_totals():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render_xlsx(_spec("overdue invoices as excel", INVOICES))))
    ws = wb["Details"]
    headers = [c.value for c in ws[1]]
    row = {h: c for h, c in zip(headers, ws[2])}
    assert isinstance(row["Due Date"].value, (datetime.date, datetime.datetime))
    assert isinstance(row["Amount"].value, float) and "AED" in row["Amount"].number_format
    assert ws.tables  # a real Excel table: filters and structured references
    total_row = [c.value for c in ws[ws.max_row]]
    assert total_row[0] == "Total" and total_row[headers.index("Amount")].startswith("=SUBTOTAL(109,")


def test_xlsx_percentages_are_numbers_and_statements_keep_their_own_totals():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render_xlsx(_spec("P&L report as excel", PNL))))
    ws = wb["Line Items"]
    pct = [r[2] for r in ws.iter_rows(min_row=2, values_only=False) if r[0].value == "Payroll"][0]
    assert pct.value == pytest.approx(0.244) and pct.number_format == "0.0%"
    assert all(c.value != "Total" for c in ws["A"])  # Total Revenue / Net Profit rows already there
    summary_text = [c.value for c in wb["Summary"]["A"] if c.value]
    assert "Executive summary" in summary_text and "Margin" in summary_text
    margin = next(r[1] for r in wb["Summary"].iter_rows(values_only=False) if r[0].value == "Margin")
    assert margin.value == pytest.approx(0.342) and margin.number_format == "0.0%"


# ── MD ───────────────────────────────────────────────────────────────────────

def test_md_carries_chart_figures_as_text():
    md = render_md(_spec("sales by customer as markdown", [
        {"customer_name": "Apex Retail", "sales": 220500}, {"customer_name": "Falcon Energy", "sales": 191310},
    ])).decode("utf-8")
    assert "| Apex Retail | 220,500.00 |" in md
